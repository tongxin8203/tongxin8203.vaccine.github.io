#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
每日文件摘要工具
扫描指定文件夹中的 PDF/PPT/Word/Excel/TXT 文件，生成摘要 PDF。
"""

import json
import os
import sys
import re
import datetime
import traceback
import subprocess


def check_and_install_dependencies():
    """检查必要的依赖库，缺失则自动安装"""
    deps = {
        "PyPDF2": "PyPDF2",
        "pptx": "python-pptx",
        "docx": "python-docx",
        "openpyxl": "openpyxl",
        "xlrd": "xlrd",
        "olefile": "olefile",
        "reportlab": "reportlab",
        "chardet": "chardet",
    }
    missing = []
    for module_name, pip_name in deps.items():
        try:
            __import__(module_name)
        except ImportError:
            missing.append(pip_name)

    if not missing:
        return

    print(f"检测到缺少依赖库: {', '.join(missing)}")
    print("正在自动安装...")
    try:
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install"] + missing,
            stdout=sys.stdout, stderr=sys.stderr
        )
        print("依赖安装完成!")
    except subprocess.CalledProcessError:
        print("=" * 50)
        print("自动安装失败，请手动运行:")
        print(f"  {sys.executable} -m pip install {' '.join(missing)}")
        print("=" * 50)
        sys.exit(1)


# --- 文件读取模块 ---

def _check_file_readable(filepath):
    """检查文件是否可读，返回 None 表示正常，否则返回错误信息"""
    if not os.path.exists(filepath):
        return "[读取失败: 文件不存在]"
    size = os.path.getsize(filepath)
    if size == 0:
        return "[读取失败: 文件大小为 0 字节（可能是未完整下载）]"
    if size < 10:
        return "[读取失败: 文件过小，可能已损坏]"
    return None


def _try_read_as_text(filepath):
    """尝试将任意文件当作文本读取（最后的备选方案）"""
    import chardet
    try:
        with open(filepath, "rb") as f:
            raw = f.read(8192)
        detected = chardet.detect(raw)
        enc = detected.get("encoding")
        conf = detected.get("confidence", 0)
        if enc and conf > 0.5:
            with open(filepath, "rb") as f:
                text = f.read().decode(enc, errors="replace")
            # 过滤掉大量不可读字符的情况
            printable = sum(1 for c in text[:500] if c.isprintable() or c in '\n\r\t')
            if printable > len(text[:500]) * 0.5:
                return text
    except Exception:
        pass
    return None


def _try_read_as_html(filepath):
    """有些 .xls/.doc 文件实际上是 HTML 格式（网页另存为）"""
    try:
        with open(filepath, "rb") as f:
            head = f.read(512)
        # 检查是否包含 HTML 标签
        head_text = head.decode("utf-8", errors="ignore").lower()
        if "<html" in head_text or "<table" in head_text or "<!doctype" in head_text:
            import chardet
            with open(filepath, "rb") as f:
                raw = f.read()
            detected = chardet.detect(raw)
            enc = detected.get("encoding", "utf-8") or "utf-8"
            text = raw.decode(enc, errors="replace")
            # 简单去除 HTML 标签
            clean = re.sub(r'<[^>]+>', ' ', text)
            clean = re.sub(r'\s+', ' ', clean).strip()
            if len(clean) > 20:
                return clean
    except Exception:
        pass
    return None


def read_pdf(filepath):
    """读取 PDF 文件文本内容"""
    err = _check_file_readable(filepath)
    if err:
        return err
    from PyPDF2 import PdfReader
    try:
        reader = PdfReader(filepath, strict=False)
        text_parts = []
        for page in reader.pages:
            try:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
            except Exception:
                continue
        if text_parts:
            return "\n".join(text_parts)
    except Exception:
        pass
    # 备选：当作文本读
    fallback = _try_read_as_text(filepath)
    if fallback:
        return fallback
    return "[读取失败: PDF 文件损坏或不完整（可能未下载完成）]"


def read_pptx(filepath):
    """读取 PPTX 文件文本内容"""
    err = _check_file_readable(filepath)
    if err:
        return err
    from pptx import Presentation
    try:
        prs = Presentation(filepath)
        text_parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            text_parts.append(text)
        return "\n".join(text_parts)
    except Exception as e:
        return f"[读取失败: {e}]"


def read_ppt_legacy(filepath):
    """读取旧版 .ppt 文件"""
    err = _check_file_readable(filepath)
    if err:
        return err
    # 先尝试 python-pptx（有些 .ppt 其实是新格式改了扩展名）
    try:
        result = read_pptx(filepath)
        if not result.startswith("[读取失败"):
            return result
    except Exception:
        pass
    # 使用 olefile 从二进制 PPT 中提取文本
    result = _extract_text_from_ole(filepath, "ppt")
    if not result.startswith("[读取失败") and not result.startswith("[文件内容为空"):
        return result
    # 备选：当作文本读
    fallback = _try_read_as_text(filepath)
    if fallback:
        return fallback
    return result


def read_docx(filepath):
    """读取 DOCX 文件文本内容"""
    err = _check_file_readable(filepath)
    if err:
        return err
    from docx import Document
    try:
        doc = Document(filepath)
        text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
        return "\n".join(text_parts)
    except Exception as e:
        return f"[读取失败: {e}]"


def read_doc_legacy(filepath):
    """读取旧版 .doc 文件"""
    err = _check_file_readable(filepath)
    if err:
        return err
    # 先尝试 python-docx（有些 .doc 其实是新格式改了扩展名）
    try:
        result = read_docx(filepath)
        if not result.startswith("[读取失败"):
            return result
    except Exception:
        pass
    # 试试是不是 HTML 格式（网页另存为 .doc 很常见）
    html_result = _try_read_as_html(filepath)
    if html_result:
        return html_result
    # 使用 olefile 从二进制 DOC 中提取文本
    result = _extract_text_from_ole(filepath, "doc")
    if not result.startswith("[读取失败") and not result.startswith("[文件内容为空"):
        return result
    # 备选：当作文本读
    fallback = _try_read_as_text(filepath)
    if fallback:
        return fallback
    return result


def read_xlsx(filepath):
    """读取 Excel 文件文本内容"""
    err = _check_file_readable(filepath)
    if err:
        return err
    from openpyxl import load_workbook
    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
        text_parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            text_parts.append(f"[工作表: {sheet_name}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                line = " | ".join(cells).strip()
                if line.replace("|", "").strip():
                    text_parts.append(line)
        wb.close()
        return "\n".join(text_parts)
    except Exception:
        pass
    # openpyxl 失败，尝试用 xlrd（可能是旧格式误用了 .xlsx 扩展名）
    result = read_xls_legacy(filepath)
    if not result.startswith("[读取失败"):
        return result
    # 试试是不是 HTML 格式
    html_result = _try_read_as_html(filepath)
    if html_result:
        return html_result
    # 备选：当作文本读
    fallback = _try_read_as_text(filepath)
    if fallback:
        return fallback
    return "[读取失败: 文件格式不兼容（可能是加密文件或未完整下载）]"


def read_xls_legacy(filepath):
    """读取旧版 .xls 文件（使用 xlrd）"""
    err = _check_file_readable(filepath)
    if err:
        return err
    import xlrd
    try:
        wb = xlrd.open_workbook(filepath)
        text_parts = []
        for sheet in wb.sheets():
            text_parts.append(f"[工作表: {sheet.name}]")
            for row_idx in range(sheet.nrows):
                cells = []
                for col_idx in range(sheet.ncols):
                    val = sheet.cell_value(row_idx, col_idx)
                    cells.append(str(val) if val else "")
                line = " | ".join(cells).strip()
                if line.replace("|", "").strip():
                    text_parts.append(line)
        return "\n".join(text_parts)
    except Exception:
        pass
    # 试试是不是 HTML 格式（很多网页导出的 .xls 实际上是 HTML）
    html_result = _try_read_as_html(filepath)
    if html_result:
        return html_result
    # 备选：当作文本读
    fallback = _try_read_as_text(filepath)
    if fallback:
        return fallback
    return "[读取失败: .xls 文件格式不兼容（可能是加密文件、网页格式或未完整下载）]"


def _extract_text_from_ole(filepath, file_type):
    """从旧版 Office 二进制文件 (OLE2) 中提取文本"""
    import olefile
    try:
        if not olefile.isOleFile(filepath):
            return "[读取失败: 文件不是有效的 OLE 格式]"

        ole = olefile.OleFileIO(filepath)
        text = ""

        if file_type == "doc":
            # Word .doc：文本存储在 "WordDocument" 流中
            # 尝试从 Word Document 流提取
            if ole.exists("WordDocument"):
                data = ole.openstream("WordDocument").read()
                # 尝试提取 UTF-16LE 编码的文本
                text = _extract_unicode_text(data)
            # 备用：遍历所有流提取文本
            if not text.strip():
                text = _extract_all_streams_text(ole)

        elif file_type == "ppt":
            # PowerPoint .ppt：文本通常在 "PowerPoint Document" 流中
            if ole.exists("PowerPoint Document"):
                data = ole.openstream("PowerPoint Document").read()
                text = _extract_ppt_text_records(data)
            if not text.strip():
                text = _extract_all_streams_text(ole)

        ole.close()
        return text if text.strip() else "[文件内容为空或无法提取文本]"

    except Exception as e:
        return f"[读取失败: {e}]"


def _extract_unicode_text(data):
    """从二进制数据中提取 Unicode (UTF-16LE) 文本片段"""
    # 查找连续的 UTF-16LE 字符序列（中文和ASCII混合）
    results = []
    try:
        decoded = data.decode("utf-16-le", errors="ignore")
        # 提取可读文本片段（至少4个字符）
        chunks = re.findall(r'[\u4e00-\u9fff\u3000-\u303f\uff00-\uffefa-zA-Z0-9\s，。！？、；：\u201c\u201d\u2018\u2019（）《》\-.,;:!?()\[\]+=/@#$%&*\\]{4,}', decoded)
        for chunk in chunks:
            cleaned = chunk.strip()
            if cleaned and len(cleaned) >= 4:
                results.append(cleaned)
    except Exception:
        pass
    return "\n".join(results)


def _extract_ppt_text_records(data):
    """从 PPT 二进制数据中提取文本记录"""
    # PPT 文本记录类型：
    # 0x0FA0 (4000) = TextCharsAtom (UTF-16LE)
    # 0x0FA8 (4008) = TextBytesAtom (ASCII/Latin-1)
    texts = []
    pos = 0
    while pos < len(data) - 8:
        # 每个记录: 2字节版本/实例, 2字节类型, 4字节长度
        rec_type = int.from_bytes(data[pos + 2:pos + 4], 'little')
        rec_len = int.from_bytes(data[pos + 4:pos + 8], 'little')

        if rec_len < 0 or rec_len > len(data) - pos - 8:
            pos += 1
            continue

        if rec_type == 0x0FA0:  # TextCharsAtom - UTF-16LE
            try:
                text = data[pos + 8:pos + 8 + rec_len].decode('utf-16-le', errors='ignore').strip()
                if text and len(text) >= 2:
                    texts.append(text)
            except Exception:
                pass
            pos += 8 + rec_len
        elif rec_type == 0x0FA8:  # TextBytesAtom - ASCII
            try:
                text = data[pos + 8:pos + 8 + rec_len].decode('latin-1', errors='ignore').strip()
                if text and len(text) >= 2:
                    texts.append(text)
            except Exception:
                pass
            pos += 8 + rec_len
        else:
            pos += 1

    return "\n".join(texts)


def _extract_all_streams_text(ole):
    """从 OLE 文件所有流中提取可读文本（备用方案）"""
    results = []
    for stream_path in ole.listdir():
        try:
            data = ole.openstream(stream_path).read()
            # 尝试 UTF-16LE
            text = _extract_unicode_text(data)
            if text.strip():
                results.append(text.strip())
        except Exception:
            continue
    return "\n".join(results)


def read_txt(filepath):
    """读取 TXT 文件文本内容，自动检测编码"""
    import chardet
    try:
        with open(filepath, "rb") as f:
            raw = f.read()
        detected = chardet.detect(raw)
        encoding = detected.get("encoding", "utf-8") or "utf-8"
        return raw.decode(encoding, errors="replace")
    except Exception as e:
        return f"[读取失败: {e}]"


# 文件类型 -> 读取函数映射
READERS = {
    ".pdf": read_pdf,
    ".pptx": read_pptx,
    ".ppt": read_ppt_legacy,
    ".docx": read_docx,
    ".doc": read_doc_legacy,
    ".xlsx": read_xlsx,
    ".xls": read_xls_legacy,
    ".txt": read_txt,
}


def make_summary(text, max_length=500):
    """简单摘要：取前 max_length 个字符"""
    text = text.strip()
    if not text:
        return "[文件内容为空]"
    if len(text) <= max_length:
        return text
    return text[:max_length] + "..."


def scan_folder(folder, file_types, date_from=None, date_to=None):
    """扫描文件夹，返回匹配的文件列表。
    date_from: 如果指定，只返回修改时间 >= date_from 00:00 的文件。
    date_to: 如果指定，只返回修改时间 < date_to+1天 00:00 的文件。
    """
    ts_from = None
    ts_to = None
    if date_from is not None:
        ts_from = datetime.datetime.combine(date_from, datetime.time.min).timestamp()
    if date_to is not None:
        next_day = date_to + datetime.timedelta(days=1)
        ts_to = datetime.datetime.combine(next_day, datetime.time.min).timestamp()

    matched = []
    for root, _dirs, files in os.walk(folder):
        for fname in files:
            ext = os.path.splitext(fname)[1].lower()
            if ext not in file_types:
                continue
            fpath = os.path.join(root, fname)
            try:
                mtime = os.path.getmtime(fpath)
            except OSError:
                continue
            if ts_from is not None and mtime < ts_from:
                continue
            if ts_to is not None and mtime >= ts_to:
                continue
            matched.append(fpath)
    matched.sort()
    return matched


def _choose_date_range(today):
    """交互式选择扫描日期范围，返回 (date_from, date_to, 描述文本)"""
    print("=" * 50)
    print("请选择要扫描的文件日期范围：")
    print("  1. 仅今天")
    print("  2. 最近 3 天")
    print("  3. 最近 7 天")
    print("  4. 本月")
    print("  5. 自定义日期")
    print("  6. 全部文件（不限日期）")
    print("=" * 50)

    while True:
        choice = input("请输入选项 [1-6]（默认 1）: ").strip()
        if choice == "" or choice == "1":
            return today, today, "仅今天"
        elif choice == "2":
            return today - datetime.timedelta(days=2), today, "最近 3 天"
        elif choice == "3":
            return today - datetime.timedelta(days=6), today, "最近 7 天"
        elif choice == "4":
            return today.replace(day=1), today, "本月"
        elif choice == "5":
            return _input_custom_date(today)
        elif choice == "6":
            return None, None, "全部文件"
        else:
            print("  无效选项，请重新输入。")


def _input_custom_date(today):
    """让用户输入自定义日期范围"""
    print()
    print("请输入日期，格式: YYYY-MM-DD")
    while True:
        start_str = input(f"  开始日期（默认 {today}）: ").strip()
        if not start_str:
            date_from = today
            break
        try:
            date_from = datetime.datetime.strptime(start_str, "%Y-%m-%d").date()
            break
        except ValueError:
            print("  日期格式不正确，请使用 YYYY-MM-DD 格式。")

    while True:
        end_str = input(f"  结束日期（默认 {today}）: ").strip()
        if not end_str:
            date_to = today
            break
        try:
            date_to = datetime.datetime.strptime(end_str, "%Y-%m-%d").date()
            break
        except ValueError:
            print("  日期格式不正确，请使用 YYYY-MM-DD 格式。")

    if date_from > date_to:
        date_from, date_to = date_to, date_from

    return date_from, date_to, f"自定义: {date_from} 至 {date_to}"


def _parse_date_arg(date_arg, today):
    """解析命令行 --date 参数，返回 (date_from, date_to, 描述文本)"""
    arg = date_arg.lower()
    if arg == "today":
        return today, today, "仅今天"
    elif arg == "3d":
        return today - datetime.timedelta(days=2), today, "最近 3 天"
    elif arg == "week":
        return today - datetime.timedelta(days=6), today, "最近 7 天"
    elif arg == "month":
        return today.replace(day=1), today, "本月"
    elif arg == "all":
        return None, None, "全部文件"
    else:
        # 尝试解析为具体日期 YYYY-MM-DD 或范围 YYYY-MM-DD~YYYY-MM-DD
        if "~" in arg:
            parts = arg.split("~", 1)
            try:
                d1 = datetime.datetime.strptime(parts[0], "%Y-%m-%d").date()
                d2 = datetime.datetime.strptime(parts[1], "%Y-%m-%d").date()
                if d1 > d2:
                    d1, d2 = d2, d1
                return d1, d2, f"自定义: {d1} 至 {d2}"
            except ValueError:
                pass
        else:
            try:
                d = datetime.datetime.strptime(arg, "%Y-%m-%d").date()
                return d, d, f"指定日期: {d}"
            except ValueError:
                pass
        print(f"警告: 无法解析日期参数 '{date_arg}'，将使用交互式选择。")
        return None  # 返回 None 表示需要交互选择



def generate_summary_pdf(summaries, output_path, date_str):
    """生成汇总 PDF 文件"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # 尝试注册中文字体
    font_name = "Helvetica"
    chinese_font_paths = [
        ("SimSun", "C:\\Windows\\Fonts\\simsun.ttc"),
        ("SimHei", "C:\\Windows\\Fonts\\simhei.ttf"),
        ("MicrosoftYaHei", "C:\\Windows\\Fonts\\msyh.ttc"),
        ("NotoSansSC", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        ("NotoSansSC", "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc"),
    ]
    for fname, fpath in chinese_font_paths:
        if os.path.exists(fpath):
            try:
                pdfmetrics.registerFont(TTFont(fname, fpath))
                font_name = fname
                break
            except Exception:
                continue

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    title_style = ParagraphStyle(
        "TitleStyle",
        fontName=font_name,
        fontSize=18,
        leading=24,
        spaceAfter=20,
        alignment=1,  # 居中
    )
    heading_style = ParagraphStyle(
        "HeadingStyle",
        fontName=font_name,
        fontSize=12,
        leading=16,
        spaceAfter=6,
        textColor="darkblue",
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        fontName=font_name,
        fontSize=10,
        leading=14,
        spaceAfter=10,
    )
    meta_style = ParagraphStyle(
        "MetaStyle",
        fontName=font_name,
        fontSize=9,
        leading=12,
        textColor="gray",
        spaceAfter=4,
    )

    elements = []

    # 标题
    title_text = f"每日文件摘要 - {date_str}"
    elements.append(Paragraph(title_text, title_style))
    elements.append(Spacer(1, 0.5 * cm))

    total_info = f"共扫描到 {len(summaries)} 个文件"
    elements.append(Paragraph(total_info, meta_style))
    elements.append(Spacer(1, 0.5 * cm))

    for i, item in enumerate(summaries, 1):
        # 文件标题
        safe_name = item["filename"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        elements.append(Paragraph(f"{i}. {safe_name}", heading_style))

        # 文件路径和类型
        safe_path = item["filepath"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        elements.append(Paragraph(f"路径: {safe_path}", meta_style))
        elements.append(Paragraph(f"类型: {item['filetype']}  大小: {item['filesize']}", meta_style))

        # 摘要内容
        safe_summary = item["summary"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
        elements.append(Paragraph(safe_summary, body_style))
        elements.append(Spacer(1, 0.3 * cm))

    doc.build(elements)


def format_size(size_bytes):
    """格式化文件大小"""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.1f} MB"


def main():
    # 检查并自动安装依赖
    check_and_install_dependencies()

    # 加载配置
    script_dir = os.path.dirname(os.path.abspath(__file__))
    config_path = os.path.join(script_dir, "config.json")

    if not os.path.exists(config_path):
        print(f"配置文件不存在: {config_path}")
        sys.exit(1)

    with open(config_path, "r", encoding="utf-8") as f:
        config = json.load(f)

    scan_folder_path = config.get("scan_folder", "")
    output_folder = config.get("output_folder", "D:\\wechat-summary")
    file_types = config.get("file_types", [".pdf", ".pptx", ".docx", ".xlsx", ".txt"])
    max_summary_length = config.get("max_summary_length", 500)

    # 解析命令行参数
    date_arg = None
    positional_args = []
    i = 1
    while i < len(sys.argv):
        if sys.argv[i] == "--date" and i + 1 < len(sys.argv):
            date_arg = sys.argv[i + 1]
            i += 2
        else:
            positional_args.append(sys.argv[i])
            i += 1

    if len(positional_args) > 0:
        scan_folder_path = positional_args[0]
    if len(positional_args) > 1:
        output_folder = positional_args[1]

    if not scan_folder_path or not os.path.isdir(scan_folder_path):
        print(f"扫描文件夹不存在: {scan_folder_path}")
        print("请修改 config.json 中的 scan_folder 或通过命令行参数指定:")
        print(f"  python {sys.argv[0]} [--date today|3d|week|month|all|YYYY-MM-DD] <扫描路径> [输出路径]")
        sys.exit(1)

    # 确保输出目录存在
    os.makedirs(output_folder, exist_ok=True)

    # 日期
    today = datetime.date.today()
    date_str = today.strftime("%Y-%m-%d")

    # 确定扫描日期范围
    if date_arg:
        # 命令行指定了 --date 参数（适合定时任务静默运行）
        result = _parse_date_arg(date_arg, today)
        if result is None:
            # 解析失败，回退到交互选择
            date_from, date_to, scan_mode = _choose_date_range(today)
        else:
            date_from, date_to, scan_mode = result
    elif sys.stdin.isatty():
        # 交互式终端，显示选择菜单
        date_from, date_to, scan_mode = _choose_date_range(today)
    else:
        # 非交互式（定时任务等），默认只扫描今天
        date_from, date_to, scan_mode = today, today, "仅今天（自动模式）"

    print()
    print(f"扫描文件夹: {scan_folder_path}")
    print(f"输出文件夹: {output_folder}")
    print(f"日期: {date_str}")
    if date_from and date_to:
        if date_from == date_to:
            print(f"扫描范围: {scan_mode}（文件修改时间: {date_from}）")
        else:
            print(f"扫描范围: {scan_mode}（文件修改时间: {date_from} 至 {date_to}）")
    else:
        print(f"扫描范围: {scan_mode}")
    print()

    # 扫描文件
    file_types_lower = [t.lower() for t in file_types]
    files = scan_folder(scan_folder_path, file_types_lower, date_from=date_from, date_to=date_to)
    print(f"找到 {len(files)} 个文件")

    if not files:
        print("没有找到匹配的文件，跳过生成。")
        return

    # 逐个读取并生成摘要
    summaries = []
    success_count = 0
    fail_count = 0

    for filepath in files:
        filename = os.path.basename(filepath)
        ext = os.path.splitext(filename)[1].lower()
        reader = READERS.get(ext)
        if not reader:
            continue

        print(f"  处理: {filename} ...", end=" ")
        try:
            text = reader(filepath)
            summary = make_summary(text, max_summary_length)
        except Exception as e:
            summary = f"[处理失败: {e}]"
            traceback.print_exc()

        # 判断是否读取成功
        is_failed = summary.startswith("[读取失败") or summary.startswith("[处理失败") or summary.startswith("[文件内容为空")
        if is_failed:
            fail_count += 1
            print(f"失败")
            print(f"    原因: {summary}")
        else:
            success_count += 1
            print(f"成功")

        try:
            fsize = format_size(os.path.getsize(filepath))
        except OSError:
            fsize = "未知"

        summaries.append({
            "filename": filename,
            "filepath": filepath,
            "filetype": ext,
            "filesize": fsize,
            "summary": summary,
        })

    # 显示统计
    print()
    print(f"处理完成: 成功 {success_count} 个, 失败 {fail_count} 个, 共 {success_count + fail_count} 个")
    if fail_count > 0:
        print("提示: 失败的文件通常是微信未完整下载、加密文件或网页另存格式，可忽略。")

    # 生成 PDF（文件被占用时自动换名）
    output_path = os.path.join(output_folder, f"{date_str}.pdf")
    for suffix in range(20):
        if suffix > 0:
            output_path = os.path.join(output_folder, f"{date_str}_{suffix}.pdf")
        try:
            # 测试文件是否可写
            with open(output_path, "ab") as _:
                pass
            break
        except PermissionError:
            print(f"  文件被占用: {output_path}，尝试换名...")
            continue

    print(f"生成摘要 PDF: {output_path}")
    generate_summary_pdf(summaries, output_path, date_str)

    print("完成!")


if __name__ == "__main__":
    main()
